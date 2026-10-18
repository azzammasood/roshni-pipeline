"""IESCO monthly maintenance schedules (owner approval 2026-10-18): one workbook per circle, plus
GSO (grid stations) and construction shutdowns, from https://iesco.com.pk/maintenance-schedule.

Each row is one planned shutdown window for a feeder on a date, with a purpose ("Annual
Maintenance", "Construction/Development Work", ...): announcements, not load shedding and not
measurements. Layouts differ by circle, so columns are found by their header words (the header is
in row 1 or 2), and merged cells take the value Excel records for the merge, which is not a guess.
Only the first sheet is read; the "OFF ON" sheets restate the same windows as events.
Rows that cannot be read are quarantined with a reason; anomalies are flagged, never fixed.
"""

import re
from datetime import UTC, date, datetime, time, timedelta
from pathlib import Path

import openpyxl

from roshni.parsers.common import ParseResult, Provenance, provenance
from roshni.parsers.iesco_xlsx import normalise_code

KARACHI = timedelta(hours=5)

# Header words -> field. Checked in order; the first match wins for each column.
HEADERS: list[tuple[str, re.Pattern[str]]] = [
    ("serial", re.compile(r"^s\.?\s*(r\.?|#)|^sr\b|^s\. ?#", re.I)),
    ("disco", re.compile(r"^discos?$", re.I)),
    ("circle", re.compile(r"^circle", re.I)),
    ("subdivision", re.compile(r"^(s/|sub)\s*div", re.I)),
    ("division", re.compile(r"^division", re.I)),
    ("feeder_code", re.compile(r"feeder\s*cod", re.I)),
    ("feeder_name", re.compile(r"name of feed|^feeder name|name of fe", re.I)),
    ("grid", re.compile(r"grid", re.I)),
    ("date", re.compile(r"^date|efe?ctive date", re.I)),
    ("start", re.compile(r"^(time\s*)?from$|^time$", re.I)),
    ("end", re.compile(r"^(time\s*)?to$", re.I)),
    ("purpose", re.compile(r"^purpose", re.I)),
    ("area", re.compile(r"area|effective area", re.I)),
    ("remarks", re.compile(r"^remarks", re.I)),
]
REQUIRED = ("date", "start", "end")
DATE_RES = [
    (re.compile(r"^(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})$"), "dmy"),
    (re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T]\d{1,2}:\d{2}(?::\d{2})?)?$"), "ymd"),
]
TIME_RE = re.compile(r"^(\d{1,2})[:.](\d{2})(?::\d{2})?\s*(h|hrs?)?$", re.I)


class MaintenanceWindow(Provenance):
    item: str  # listing-page row: islamabad, rawalpindi_city, ..., gso, construction
    sheet: str
    sheet_row: int  # 1-based row in the sheet
    serial: str | None
    circle_as_printed: str | None
    division: str | None
    subdivision: str | None
    grid: str | None
    feeder_name: str | None
    feeder_code_raw: str | None
    feeder_code: str | None  # 6-digit zero-padded text; None when not numeric
    outage_date: date
    start_local: str  # "HH:MM", Asia/Karachi
    end_local: str  # "HH:MM"; "24:00" is midnight at the end of the day
    start_utc: datetime
    end_utc: datetime
    purpose: str | None
    area_affected: str | None
    flags: str  # ";"-separated, empty when clean


def _clean(v: object) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def parse_date(v: object) -> tuple[date | None, time | None]:
    """A date cell -> (date, time carried by the cell); (None, None) if unreadable."""
    if isinstance(v, datetime):
        t = v.time()
        return v.date(), (t if t != time(0) else None)
    if isinstance(v, date):
        return v, None
    s = _clean(v)
    if not s:
        return None, None
    for rx, order in DATE_RES:
        m = rx.match(s)
        if m:
            a, b, c = (int(x) for x in m.groups())
            try:
                return (date(c, b, a) if order == "dmy" else date(a, b, c)), None
            except ValueError:
                return None, None
    return None, None


MULTI_DATE_RE = re.compile(r"\d{1,2}[.\-/]\d{1,2}[.\-/]\d{4}")


def parse_dates(v: object) -> list[tuple[date, time | None]]:
    """Every date in a cell. One cell may list several dates for the same window
    ("28-09-2026 01-10-2026 ..."); each listed date is a window. [] if any token is unreadable."""
    if isinstance(v, str):
        found = MULTI_DATE_RE.findall(v)
        rest = MULTI_DATE_RE.sub(" ", v).replace(",", " ").strip()
        if len(found) > 1 and not rest:
            out = [parse_date(tok) for tok in found]
            return [(d, t) for d, t in out if d is not None] if all(d for d, _ in out) else []
    d, t = parse_date(v)
    return [(d, t)] if d is not None else []


def parse_time(v: object) -> str | None:
    """A time cell -> "HH:MM" (24:00 allowed for midnight at the end). None if unreadable."""
    if isinstance(v, datetime):
        v = v.time()
    if isinstance(v, time):
        return f"{v.hour:02d}:{v.minute:02d}"
    s = _clean(v)
    if not s:
        return None
    m = TIME_RE.match(s)
    if not m:
        return None
    h, mi = int(m[1]), int(m[2])
    if not (0 <= h <= 24 and 0 <= mi < 60) or (h == 24 and mi):
        return None
    return f"{h:02d}:{mi:02d}"


def _minutes(hhmm: str) -> int:
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


def _utc(d: date, hhmm: str) -> datetime:
    local = datetime(d.year, d.month, d.day, tzinfo=UTC) + timedelta(minutes=_minutes(hhmm))
    return local - KARACHI


def _header(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    """(row index, field -> column) for the first of rows 1-3 that names the required columns.

    A lone "TIME" header spanning two columns (From/To underneath) maps start to it and end to
    the next column.
    """
    for i, row in enumerate(rows[:3]):
        cols: dict[str, int] = {}
        for j, cell in enumerate(row):
            text = _clean(cell)
            if not text:
                continue
            for name, rx in HEADERS:
                if name not in cols and rx.search(text):
                    cols[name] = j
                    break
        if "start" in cols and "end" not in cols and _clean(rows[i][cols["start"]]) == "TIME":
            cols["end"] = cols["start"] + 1
        if all(k in cols for k in REQUIRED) and ("feeder_name" in cols or "feeder_code" in cols):
            return i, cols
    return None


def _unmerged(ws) -> list[list[object]]:
    """Sheet values with each merged range filled from its top-left cell (as Excel shows it)."""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:
        top = grid[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(grid) else None
        for r in range(rng.min_row - 1, min(rng.max_row, len(grid))):
            for c in range(rng.min_col - 1, rng.max_col):
                if c < len(grid[r]):
                    grid[r][c] = top
    return grid


def parse_workbook(path: Path, item: str) -> ParseResult[MaintenanceWindow]:
    res: ParseResult[MaintenanceWindow] = ParseResult("iesco_maintenance_window")
    prov = provenance(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = _unmerged(ws)
    head = _header([tuple(r) for r in rows])
    if head is None:
        why = "no header row with date, from, to and feeder"
        res.reject(prov, f"sheet {ws.title!r}", str(rows[:2])[:500], why)
        return res
    hi, cols = head
    first = hi + 1
    # A "From/To" sub-header row directly under the header carries no data.
    sub = {_clean(c) for c in rows[first] if _clean(c)} if first < len(rows) else set()
    if sub and sub <= {"From", "To", "FROM", "TO"}:
        first += 1

    def cell(row: list[object], name: str) -> object:
        j = cols.get(name)
        return row[j] if j is not None and j < len(row) else None

    for i in range(first, len(rows)):
        row = rows[i]
        texts = [_clean(c) for c in row if _clean(c)]
        if not texts:
            continue
        loc = f"sheet {ws.title!r} row {i + 1}"
        raw = " | ".join(texts)
        dates = parse_dates(cell(row, "date"))
        start, end = parse_time(cell(row, "start")), parse_time(cell(row, "end"))
        name, code_raw = _clean(cell(row, "feeder_name")), _clean(cell(row, "feeder_code"))
        if not dates:
            res.reject(prov, loc, raw, f"date not readable: {_clean(cell(row, 'date'))!r}")
            continue
        if start is None or end is None:
            got = f"{_clean(cell(row, 'start'))!r} to {_clean(cell(row, 'end'))!r}"
            res.reject(prov, loc, raw, f"time not readable: {got}")
            continue
        if _minutes(end) <= _minutes(start):
            res.reject(prov, loc, raw, f"window does not end after it starts: {start} to {end}")
            continue
        if not name and not code_raw:
            res.reject(prov, loc, raw, "no feeder name or code")
            continue
        base_flags = ["multi_date_cell"] if len(dates) > 1 else []
        code = normalise_code(cell(row, "feeder_code")) if code_raw else None
        if code_raw and code is None:
            base_flags.append("feeder_code_not_numeric")
        circle = _clean(cell(row, "circle"))
        if circle and "/" in circle:
            base_flags.append("circle_shared")  # e.g. "ISLAMABAD/PESCO"
        for d, cell_time in dates:
            flags = list(base_flags)
            if cell_time is not None and f"{cell_time.hour:02d}:{cell_time.minute:02d}" != start:
                flags.append(f"date_cell_time_{cell_time.hour:02d}{cell_time.minute:02d}")
            res.rows.append(
                MaintenanceWindow(
                    **prov.model_dump(),
                    item=item,
                    sheet=ws.title,
                    sheet_row=i + 1,
                    serial=_clean(cell(row, "serial")),
                    circle_as_printed=circle,
                    division=_clean(cell(row, "division")),
                    subdivision=_clean(cell(row, "subdivision")),
                    grid=_clean(cell(row, "grid")),
                    feeder_name=name,
                    feeder_code_raw=code_raw,
                    feeder_code=code,
                    outage_date=d,
                    start_local=start,
                    end_local=end,
                    start_utc=_utc(d, start),
                    end_utc=_utc(d, end),
                    purpose=_clean(cell(row, "purpose")),
                    area_affected=_clean(cell(row, "area")),
                    flags=";".join(flags),
                )
            )
    return res
